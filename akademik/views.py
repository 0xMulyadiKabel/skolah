# Simpan sebagai: akademik/views.py

from datetime import date

from django.contrib import messages
from django.contrib.auth import get_user_model
from django.core.paginator import Paginator
from django.db.models import Q
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render

from absensi.forms import KoreksiAbsensiForm
from absensi.models import AbsensiHarian
from absensi.utils import catat_koreksi_absensi
from accounts.utils import get_school
from accounts.views import admin_required

from .forms import GuruAccountForm, GuruEditForm, ImportGuruForm, ImportSiswaForm, KelasForm, OrangTuaAccountForm, OrangTuaEditForm, SiswaForm
from .models import Guru, Kelas, OrangTua, Siswa

User = get_user_model()


def _school_or_guard(request):
    """Kembalikan (school, None) kalau ada, atau (None, response_guard) kalau
    School belum ada sama sekali -- panggil di awal tiap view, langsung
    `return guard` kalau guard bukan None."""
    school = get_school(request)
    if school is None:
        return None, render(request, "accounts/school_missing.html", {"page_title": "Data Siswa"})
    return school, None


# ================= SISWA =================

@admin_required
def siswa_list(request):
    school, guard = _school_or_guard(request)
    if guard:
        return guard

    q = request.GET.get("q", "").strip()
    kelas_id = request.GET.get("kelas", "")

    siswa_qs = Siswa.objects.filter(school=school).select_related("kelas", "user").order_by("nama")
    if q:
        siswa_qs = siswa_qs.filter(Q(nama__icontains=q) | Q(nis__icontains=q))
    if kelas_id:
        siswa_qs = siswa_qs.filter(kelas_id=kelas_id)

    paginator = Paginator(siswa_qs, 20)
    page_obj = paginator.get_page(request.GET.get("page"))

    context = {
        "page_title": "Data Siswa",
        "page_obj": page_obj,
        "kelas_list": Kelas.objects.filter(school=school).order_by("nama_kelas"),
        "q": q,
        "kelas_id_selected": str(kelas_id),
    }
    return render(request, "akademik/siswa_list.html", context)


@admin_required
def siswa_create(request):
    school, guard = _school_or_guard(request)
    if guard:
        return guard

    if request.method == "POST":
        form = SiswaForm(request.POST, school=school)
        if form.is_valid():
            siswa = form.save(commit=False)
            siswa.school = school
            siswa.save()

            pesan = f"Siswa {siswa.nama} berhasil ditambahkan."
            if form.cleaned_data.get("buat_akun_sekaligus"):
                if not siswa.nis:
                    pesan += " Akun login BELUM dibuat karena NIS kosong (isi NIS dulu, lalu buat akun manual dari Data Siswa)."
                else:
                    from django.utils.crypto import get_random_string
                    temp_password = get_random_string(8)
                    user = User.objects.create_user(
                        username=siswa.nis, password=temp_password, role=User.Role.SISWA, school=school,
                    )
                    siswa.user = user
                    siswa.save()
                    pesan += f" Akun login sekaligus dibuat — Username: {siswa.nis}, Password: {temp_password}."

            messages.success(request, pesan)
            return redirect("akademik:siswa_list")
    else:
        form = SiswaForm(school=school)
    return render(request, "akademik/siswa_form.html", {"page_title": "Tambah Siswa", "form": form})


@admin_required
def siswa_update(request, pk):
    school, guard = _school_or_guard(request)
    if guard:
        return guard

    siswa = get_object_or_404(Siswa, pk=pk, school=school)
    if request.method == "POST":
        form = SiswaForm(request.POST, instance=siswa, school=school)
        if form.is_valid():
            form.save()
            messages.success(request, f"Data {siswa.nama} berhasil diperbarui.")
            return redirect("akademik:siswa_list")
    else:
        form = SiswaForm(instance=siswa, school=school)
    return render(request, "akademik/siswa_form.html", {"page_title": f"Edit Siswa — {siswa.nama}", "form": form, "siswa": siswa})


@admin_required
def siswa_kartu(request, pk):
    school, guard = _school_or_guard(request)
    if guard:
        return guard
    siswa = get_object_or_404(Siswa, pk=pk, school=school)
    if not siswa.nis:
        messages.error(request, f"{siswa.nama} belum punya NIS. Isi NIS dulu sebelum cetak kartu (QR kartu berisi NIS).")
        return redirect("akademik:siswa_update", pk=siswa.pk)
    return render(request, "akademik/siswa_kartu.html", {"page_title": f"Kartu — {siswa.nama}", "siswa": siswa})


@admin_required
def siswa_kartu_qr_image(request, pk):
    import io

    import qrcode

    school = get_school(request)
    siswa = get_object_or_404(Siswa, pk=pk, school=school)

    img = qrcode.make(siswa.nis)
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return HttpResponse(buf.getvalue(), content_type="image/png")


def _normalisasi_nis(value):
    """Excel kadang nyimpen NIS sebagai angka (float), bukan teks -- ini
    yang bikin '56781234' bisa kebaca jadi '56781234.0' kalau tidak
    dirapikan dulu. Fungsi ini nyamain semua jadi teks bersih."""
    if value is None:
        return ""
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return str(value)
    return str(value).strip()


@admin_required
def siswa_import_template(request):
    import io

    import openpyxl
    from openpyxl.styles import Font, PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data Siswa"

    headers = ["NIS", "Nama", "Kelas", "Jenis Kelamin (L/P)"]
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1B5E45")

    contoh = ["11099", "Ahmad Fadillah (CONTOH -- hapus baris ini)", "X-1", "L"]
    for i, v in enumerate(contoh, start=1):
        ws.cell(row=2, column=i, value=v)

    widths = [14, 34, 14, 18]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    response = HttpResponse(
        buf.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="template_import_siswa.xlsx"'
    return response


def _proses_import_excel(file_obj, school):
    import openpyxl

    wb = openpyxl.load_workbook(file_obj, data_only=True)
    ws = wb.active

    # Cari kolom berdasarkan NAMA header (bukan urutan tetap), supaya
    # tetap jalan walau kolomnya sengaja/tidak sengaja dipindah urutannya.
    header_map = {}
    for idx, c in enumerate(ws[1], start=1):
        if not c.value:
            continue
        key = str(c.value).strip().lower()
        if "nis" in key:
            header_map["nis"] = idx
        elif "nama" in key:
            header_map["nama"] = idx
        elif "kelas" in key:
            header_map["kelas"] = idx
        elif "kelamin" in key or key == "jk":
            header_map["jk"] = idx

    kelas_map = {k.nama_kelas.strip().lower(): k for k in Kelas.objects.filter(school=school)}
    nis_terpakai = set(Siswa.objects.filter(school=school, nis__isnull=False).values_list("nis", flat=True))
    nis_dalam_file = set()

    baris_hasil = []
    total = berhasil = gagal = 0

    for row_num in range(2, ws.max_row + 1):
        nama_cell = ws.cell(row=row_num, column=header_map.get("nama", 2)).value
        if not nama_cell or not str(nama_cell).strip():
            continue  # baris kosong dilewati diam-diam, tidak dihitung error

        total += 1
        nama = str(nama_cell).strip()
        nis = _normalisasi_nis(ws.cell(row=row_num, column=header_map.get("nis", 1)).value)
        kelas_raw = ws.cell(row=row_num, column=header_map.get("kelas", 3)).value
        kelas_key = str(kelas_raw).strip().lower() if kelas_raw else ""
        jk_raw = ws.cell(row=row_num, column=header_map.get("jk", 4)).value
        jk = str(jk_raw).strip().upper() if jk_raw else ""

        errors = []
        kelas_obj = kelas_map.get(kelas_key)
        if not kelas_obj:
            errors.append(f"Kelas '{kelas_raw or '(kosong)'}' tidak ditemukan -- cek ejaan/buat dulu di halaman Kelas")
        if jk not in ("L", "P"):
            errors.append("Jenis kelamin harus diisi L atau P")
        if nis:
            if nis in nis_terpakai:
                errors.append(f"NIS {nis} sudah dipakai siswa lain")
            elif nis in nis_dalam_file:
                errors.append(f"NIS {nis} duplikat di dalam file ini")

        if errors:
            gagal += 1
            baris_hasil.append({"baris": row_num, "nama": nama, "nis": nis, "status": "gagal", "pesan": "; ".join(errors)})
            continue

        Siswa.objects.create(school=school, nis=nis or None, nama=nama, kelas=kelas_obj, jenis_kelamin=jk, aktif=True)
        if nis:
            nis_terpakai.add(nis)
            nis_dalam_file.add(nis)
        berhasil += 1
        baris_hasil.append({"baris": row_num, "nama": nama, "nis": nis, "status": "berhasil", "pesan": "Tersimpan"})

    return {"total": total, "berhasil": berhasil, "gagal": gagal, "baris": baris_hasil}


@admin_required
def siswa_import(request):
    school, guard = _school_or_guard(request)
    if guard:
        return guard

    hasil = None
    if request.method == "POST":
        form = ImportSiswaForm(request.POST, request.FILES)
        if form.is_valid():
            hasil = _proses_import_excel(request.FILES["file_excel"], school)
            if hasil["berhasil"]:
                messages.success(request, f"Import selesai: {hasil['berhasil']} siswa berhasil ditambahkan, {hasil['gagal']} baris gagal.")
            else:
                messages.warning(request, f"Import selesai, tapi tidak ada siswa yang berhasil ditambahkan ({hasil['gagal']} baris gagal). Cek detail di bawah.")
    else:
        form = ImportSiswaForm()

    return render(request, "akademik/siswa_import.html", {"page_title": "Migrasi Data Siswa", "form": form, "hasil": hasil})


@admin_required
def siswa_bulk_nonaktifkan(request):
    school, guard = _school_or_guard(request)
    if guard:
        return guard

    if request.method == "POST":
        ids = request.POST.getlist("siswa_ids")
        updated = Siswa.objects.filter(school=school, pk__in=ids).update(aktif=False)
        messages.success(request, f"{updated} siswa berhasil dinonaktifkan.")
    return redirect("akademik:siswa_list")


@admin_required
def siswa_buat_akun(request, pk):
    from django.utils.crypto import get_random_string

    school = get_school(request)
    siswa = get_object_or_404(Siswa, pk=pk, school=school)

    if request.method == "POST":
        if siswa.user_id:
            messages.warning(request, f"{siswa.nama} sudah punya akun login.")
        elif not siswa.nis:
            messages.error(request, f"{siswa.nama} belum punya NIS. Isi NIS dulu sebelum buat akun (NIS dipakai sebagai username login).")
        else:
            temp_password = get_random_string(8)
            user = User.objects.create_user(
                username=siswa.nis, password=temp_password, role=User.Role.SISWA, school=school,
            )
            siswa.user = user
            siswa.save()
            messages.success(
                request,
                f"Akun {siswa.nama} dibuat — Username: {siswa.nis}, Password: {temp_password} (sampaikan ke siswa/orang tua).",
            )
    return redirect("akademik:siswa_list")


@admin_required
def siswa_reset_password(request, pk):
    from django.utils.crypto import get_random_string

    school = get_school(request)
    siswa = get_object_or_404(Siswa, pk=pk, school=school)

    if request.method == "POST":
        if not siswa.user_id:
            messages.error(request, f"{siswa.nama} belum punya akun login, buat akunnya dulu.")
        else:
            temp_password = get_random_string(8)
            siswa.user.set_password(temp_password)
            siswa.user.save()
            messages.success(request, f"Kata sandi {siswa.nama} direset jadi: {temp_password} — sampaikan ke yang bersangkutan.")
    return redirect("akademik:siswa_list")


@admin_required
def siswa_akun_toggle(request, pk):
    """
    Toggle status AKUN LOGIN siswa (User.is_active) -- ini SENGAJA konsep
    terpisah dari Siswa.aktif (status keaktifan siswa di sekolah, dipakai
    fitur nonaktifkan-massal). Siswa bisa saja masih aktif terdaftar di
    sekolah tapi akun login-nya sementara dinonaktifkan (mis. lupa sandi
    berkali-kali, atau kasus lain), atau sebaliknya.
    """
    school = get_school(request)
    siswa = get_object_or_404(Siswa, pk=pk, school=school)

    if request.method == "POST":
        if not siswa.user_id:
            messages.error(request, f"{siswa.nama} belum punya akun login.")
        else:
            siswa.user.is_active = not siswa.user.is_active
            siswa.user.save()
            status = "diaktifkan" if siswa.user.is_active else "dinonaktifkan"
            messages.success(request, f"Akun login {siswa.nama} berhasil {status}.")
    return redirect("akademik:siswa_list")


# ================= KELAS =================

@admin_required
def kelas_list(request):
    school, guard = _school_or_guard(request)
    if guard:
        return guard
    q = request.GET.get("q", "").strip()
    kelas_qs = Kelas.objects.filter(school=school).select_related("wali_kelas").order_by("nama_kelas")
    if q:
        kelas_qs = kelas_qs.filter(Q(nama_kelas__icontains=q) | Q(wali_kelas__nama__icontains=q))
    return render(request, "akademik/kelas_list.html", {"page_title": "Kelas", "kelas_list": kelas_qs, "q": q})


@admin_required
def kelas_create(request):
    school, guard = _school_or_guard(request)
    if guard:
        return guard

    if request.method == "POST":
        form = KelasForm(request.POST, school=school)
        if form.is_valid():
            kelas = form.save(commit=False)
            kelas.school = school
            kelas.save()
            messages.success(request, f"Kelas {kelas.nama_kelas} berhasil ditambahkan.")
            return redirect("akademik:kelas_list")
    else:
        form = KelasForm(school=school)
    return render(request, "akademik/kelas_form.html", {"page_title": "Tambah Kelas", "form": form})


@admin_required
def kelas_update(request, pk):
    school, guard = _school_or_guard(request)
    if guard:
        return guard

    kelas = get_object_or_404(Kelas, pk=pk, school=school)
    if request.method == "POST":
        form = KelasForm(request.POST, instance=kelas, school=school)
        if form.is_valid():
            form.save()
            messages.success(request, f"Kelas {kelas.nama_kelas} berhasil diperbarui.")
            return redirect("akademik:kelas_list")
    else:
        form = KelasForm(instance=kelas, school=school)
    return render(request, "akademik/kelas_form.html", {"page_title": f"Edit Kelas — {kelas.nama_kelas}", "form": form})


# ================= KOREKSI ABSENSI (ADMIN) =================

@admin_required
def koreksi_kelas_list(request, kelas_id):
    school = get_school(request)
    kelas = get_object_or_404(Kelas, pk=kelas_id, school=school)
    today = date.today()

    siswa_list = Siswa.objects.filter(kelas=kelas, aktif=True).order_by("nama")
    absensi_map = {a.siswa_id: a for a in AbsensiHarian.objects.filter(siswa__kelas=kelas, tanggal=today)}
    baris = [{"siswa": s, "absensi": absensi_map.get(s.id)} for s in siswa_list]

    context = {"page_title": f"Koreksi Absensi \u2014 {kelas.nama_kelas}", "kelas": kelas, "baris": baris, "tanggal": today}
    return render(request, "akademik/koreksi_kelas_list.html", context)


@admin_required
def koreksi_absensi(request, siswa_id):
    school = get_school(request)
    siswa = get_object_or_404(Siswa, pk=siswa_id, school=school)

    tanggal_str = request.GET.get("tanggal")
    tanggal = date.fromisoformat(tanggal_str) if tanggal_str else date.today()
    absensi_ada = AbsensiHarian.objects.filter(siswa=siswa, tanggal=tanggal).first()

    if request.method == "POST":
        form = KoreksiAbsensiForm(request.POST)
        if form.is_valid():
            catat_koreksi_absensi(
                siswa=siswa, tanggal=tanggal,
                status=form.cleaned_data["status"],
                jam_masuk=form.cleaned_data["jam_masuk"],
                jam_pulang=form.cleaned_data["jam_pulang"],
                catatan=form.cleaned_data["catatan"],
                user=request.user,
            )
            messages.success(request, f"Kehadiran {siswa.nama} tanggal {tanggal} berhasil dikoreksi.")
            return redirect("akademik:koreksi_kelas_list", kelas_id=siswa.kelas_id)
    else:
        initial = {}
        if absensi_ada:
            initial = {"status": absensi_ada.status, "jam_masuk": absensi_ada.jam_masuk, "jam_pulang": absensi_ada.jam_pulang}
        form = KoreksiAbsensiForm(initial=initial)

    context = {
        "page_title": f"Koreksi Absensi \u2014 {siswa.nama}",
        "form": form, "siswa": siswa, "tanggal": tanggal, "absensi_ada": absensi_ada,
    }
    return render(request, "akademik/koreksi_absensi_form.html", context)


# ================= AKUN GURU =================

@admin_required
def akun_guru_list(request):
    school, guard = _school_or_guard(request)
    if guard:
        return guard
    guru_list = Guru.objects.filter(school=school).select_related("user").prefetch_related("kelas_diampu")
    return render(request, "akademik/akun_guru_list.html", {"page_title": "Akun Guru", "guru_list": guru_list})


@admin_required
def akun_guru_create(request):
    school, guard = _school_or_guard(request)
    if guard:
        return guard

    if request.method == "POST":
        form = GuruAccountForm(request.POST)
        if form.is_valid():
            user = User.objects.create_user(
                username=form.cleaned_data["username"],
                password=form.cleaned_data["password"],
                role=User.Role.GURU,
                school=school,
            )
            Guru.objects.create(school=school, user=user, nama=form.cleaned_data["nama"], email=form.cleaned_data["email"])
            messages.success(request, f"Akun guru {form.cleaned_data['nama']} berhasil dibuat.")
            return redirect("akademik:akun_guru_list")
    else:
        form = GuruAccountForm()
    return render(request, "akademik/akun_form.html", {"page_title": "Tambah Akun Guru", "form": form})


@admin_required
def akun_guru_edit(request, pk):
    guru = get_object_or_404(Guru, pk=pk, school=get_school(request))
    if request.method == "POST":
        form = GuruEditForm(request.POST, instance=guru)
        if form.is_valid():
            form.save()
            messages.success(request, f"Data {guru.nama} berhasil diperbarui.")
            return redirect("akademik:akun_guru_list")
    else:
        form = GuruEditForm(instance=guru)
    return render(request, "akademik/akun_form.html", {"page_title": f"Edit Akun Guru — {guru.nama}", "form": form})


@admin_required
def akun_guru_reset(request, pk):
    from django.utils.crypto import get_random_string
    guru = get_object_or_404(Guru, pk=pk, school=get_school(request))
    if request.method == "POST":
        temp_password = get_random_string(8)
        guru.user.set_password(temp_password)
        guru.user.save()
        messages.success(request, f"Kata sandi {guru.nama} direset jadi: {temp_password} — sampaikan ke yang bersangkutan.")
    return redirect("akademik:akun_guru_list")


@admin_required
def akun_guru_toggle_active(request, pk):
    guru = get_object_or_404(Guru, pk=pk, school=get_school(request))
    if request.method == "POST":
        guru.user.is_active = not guru.user.is_active
        guru.user.save()
        status = "diaktifkan" if guru.user.is_active else "dinonaktifkan"
        messages.success(request, f"Akun {guru.nama} berhasil {status}.")
    return redirect("akademik:akun_guru_list")


# ================= AKUN ORANG TUA =================

@admin_required
def akun_ortu_list(request):
    school, guard = _school_or_guard(request)
    if guard:
        return guard
    ortu_list = OrangTua.objects.filter(school=school).select_related("user").prefetch_related("anak")
    return render(request, "akademik/akun_ortu_list.html", {"page_title": "Akun Orang Tua", "ortu_list": ortu_list})


@admin_required
def akun_ortu_create(request):
    school, guard = _school_or_guard(request)
    if guard:
        return guard

    if request.method == "POST":
        form = OrangTuaAccountForm(request.POST, school=school)
        if form.is_valid():
            user = User.objects.create_user(
                username=form.cleaned_data["username"],
                password=form.cleaned_data["password"],
                role=User.Role.ORANG_TUA,
                school=school,
            )
            ortu = OrangTua.objects.create(school=school, user=user, nama=form.cleaned_data["nama"], no_hp=form.cleaned_data["no_hp"])
            ortu.anak.set(form.cleaned_data["anak"])
            messages.success(request, f"Akun orang tua {form.cleaned_data['nama']} berhasil dibuat.")
            return redirect("akademik:akun_ortu_list")
    else:
        form = OrangTuaAccountForm(school=school)
    return render(request, "akademik/akun_form.html", {"page_title": "Tambah Akun Orang Tua", "form": form})


@admin_required
def akun_ortu_edit(request, pk):
    school = get_school(request)
    ortu = get_object_or_404(OrangTua, pk=pk, school=school)
    if request.method == "POST":
        form = OrangTuaEditForm(request.POST, instance=ortu, school=school)
        if form.is_valid():
            form.save()
            messages.success(request, f"Data {ortu.nama} berhasil diperbarui.")
            return redirect("akademik:akun_ortu_list")
    else:
        form = OrangTuaEditForm(instance=ortu, school=school)
    return render(request, "akademik/akun_form.html", {"page_title": f"Edit Akun Orang Tua — {ortu.nama}", "form": form})


@admin_required
def akun_ortu_reset(request, pk):
    from django.utils.crypto import get_random_string
    ortu = get_object_or_404(OrangTua, pk=pk, school=get_school(request))
    if request.method == "POST":
        temp_password = get_random_string(8)
        ortu.user.set_password(temp_password)
        ortu.user.save()
        messages.success(request, f"Kata sandi {ortu.nama} direset jadi: {temp_password} — sampaikan ke yang bersangkutan.")
    return redirect("akademik:akun_ortu_list")


@admin_required
def akun_ortu_toggle_active(request, pk):
    ortu = get_object_or_404(OrangTua, pk=pk, school=get_school(request))
    if request.method == "POST":
        ortu.user.is_active = not ortu.user.is_active
        ortu.user.save()
        status = "diaktifkan" if ortu.user.is_active else "dinonaktifkan"
        messages.success(request, f"Akun {ortu.nama} berhasil {status}.")
    return redirect("akademik:akun_ortu_list")


# ================= IMPORT DATA GURU (EXCEL) =================

def _proses_import_guru_excel(file_obj, school):
    import openpyxl

    wb = openpyxl.load_workbook(file_obj, data_only=True)
    ws = wb.active

    header_map = {}
    for idx, c in enumerate(ws[1], start=1):
        if not c.value:
            continue
        key = str(c.value).strip().lower()
        if "username" in key:
            header_map["username"] = idx
        elif "nama" in key:
            header_map["nama"] = idx
        elif "email" in key:
            header_map["email"] = idx

    username_terpakai = set(User.objects.values_list("username", flat=True))  # global, username unik system-wide
    username_dalam_file = set()

    baris_hasil = []
    total = berhasil = gagal = 0

    for row_num in range(2, ws.max_row + 1):
        nama_cell = ws.cell(row=row_num, column=header_map.get("nama", 2)).value
        if not nama_cell or not str(nama_cell).strip():
            continue

        total += 1
        nama = str(nama_cell).strip()
        username_raw = ws.cell(row=row_num, column=header_map.get("username", 1)).value
        username = str(username_raw).strip() if username_raw else ""
        email_raw = ws.cell(row=row_num, column=header_map.get("email", 3)).value
        email = str(email_raw).strip() if email_raw else ""

        errors = []
        if not username:
            errors.append("Username kosong")
        elif username in username_terpakai:
            errors.append(f"Username '{username}' sudah dipakai")
        elif username in username_dalam_file:
            errors.append(f"Username '{username}' duplikat di dalam file ini")
        if not email:
            errors.append("Email kosong")

        if errors:
            gagal += 1
            baris_hasil.append({"baris": row_num, "nama": nama, "username": username, "status": "gagal", "pesan": "; ".join(errors)})
            continue

        from django.utils.crypto import get_random_string
        temp_password = get_random_string(8)
        user = User.objects.create_user(username=username, password=temp_password, role=User.Role.GURU, school=school)
        Guru.objects.create(school=school, user=user, nama=nama, email=email)

        username_terpakai.add(username)
        username_dalam_file.add(username)
        berhasil += 1
        baris_hasil.append({
            "baris": row_num, "nama": nama, "username": username, "status": "berhasil",
            "pesan": f"Tersimpan. Password sementara: {temp_password}",
        })

    return {"total": total, "berhasil": berhasil, "gagal": gagal, "baris": baris_hasil}


@admin_required
def guru_import_template(request):
    import io

    import openpyxl
    from openpyxl.styles import Font, PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data Guru"

    headers = ["Username", "Nama", "Email"]
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1B5E45")

    contoh = ["siti.aminah", "Siti Aminah, S.Pd (CONTOH -- hapus baris ini)", "siti.aminah@sekolah.sch.id"]
    for i, v in enumerate(contoh, start=1):
        ws.cell(row=2, column=i, value=v)

    for i, w in enumerate([18, 34, 28], start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    response = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = 'attachment; filename="template_import_guru.xlsx"'
    return response


@admin_required
def guru_kartu(request, pk):
    school = get_school(request)
    guru = get_object_or_404(Guru, pk=pk, school=school)
    return render(request, "akademik/guru_kartu.html", {"page_title": f"Kartu — {guru.nama}", "guru": guru})


@admin_required
def guru_kartu_qr_image(request, pk):
    import io

    import qrcode

    school = get_school(request)
    guru = get_object_or_404(Guru, pk=pk, school=school)

    img = qrcode.make(guru.user.username)
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return HttpResponse(buf.getvalue(), content_type="image/png")


@admin_required
def guru_import(request):
    school = get_school(request)
    hasil = None
    if request.method == "POST":
        form = ImportGuruForm(request.POST, request.FILES)
        if form.is_valid():
            hasil = _proses_import_guru_excel(request.FILES["file_excel"], school)
            if hasil["berhasil"]:
                messages.success(request, f"Import selesai: {hasil['berhasil']} akun guru berhasil dibuat, {hasil['gagal']} baris gagal.")
            else:
                messages.warning(request, f"Import selesai, tidak ada akun yang berhasil dibuat ({hasil['gagal']} baris gagal).")
    else:
        form = ImportGuruForm()
    return render(request, "akademik/guru_import.html", {"page_title": "Migrasi Data Guru", "form": form, "hasil": hasil})